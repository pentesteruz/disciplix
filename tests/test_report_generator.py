"""
Excel/PDF hisobotlari.

Excel yozish pandas'dan openpyxl'ga ko'chirildi — pandas bitta jadval uchun
~100 MB bog'liqlik edi, openpyxl esa allaqachon talab qilinardi.
Bu testlar ko'chirish natijani o'zgartirmaganini tekshiradi.
"""

import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from utils.report_generator import generate_excel_report, generate_pdf_report


class FakeTx:
    def __init__(self, amount, tx_type="expense", category="Oziq-ovqat",
                 currency="UZS", description="Tushlik", timestamp=None):
        self.amount = amount
        self.type = tx_type
        self.category = category
        self.currency = currency
        self.description = description
        self.timestamp = timestamp or datetime(2026, 8, 1, 12, 30)


class FakeTask:
    def __init__(self, title="Sport", is_completed=False, due_date=None):
        self.title = title
        self.is_completed = is_completed
        self.due_date = due_date


@pytest.mark.asyncio
async def test_excel_ikkala_varaqni_yaratadi():
    out = await generate_excel_report(
        [FakeTx(20000)], [FakeTask()], datetime(2026, 8, 1), datetime(2026, 8, 3)
    )
    wb = load_workbook(out)
    assert wb.sheetnames == ["Tranzaksiyalar", "Rejalar"]


@pytest.mark.asyncio
async def test_excel_tranzaksiya_qatorlari():
    txs = [FakeTx(20000, "expense", "Oziq-ovqat"), FakeTx(500000, "income", "Oylik")]
    out = await generate_excel_report(txs, [], datetime(2026, 8, 1), datetime(2026, 8, 3))

    sheet = load_workbook(out)["Tranzaksiyalar"]
    sarlavhalar = [c.value for c in sheet[1]]
    assert sarlavhalar == ["Sana", "Turi", "Kategoriya", "Miqdor", "Valyuta", "Izoh"]

    assert sheet.cell(row=2, column=2).value == "Chiqim"
    assert sheet.cell(row=2, column=4).value == 20000
    assert sheet.cell(row=3, column=2).value == "Kirim"


@pytest.mark.asyncio
async def test_excel_bosh_royxatda_xabar_yozadi():
    out = await generate_excel_report([], [], datetime(2026, 8, 1), datetime(2026, 8, 3))
    wb = load_workbook(out)
    assert wb["Tranzaksiyalar"].cell(row=2, column=1).value == "Ushbu davrda tranzaksiyalar yo'q"
    assert wb["Rejalar"].cell(row=2, column=1).value == "Ushbu davrda vazifalar yo'q"


@pytest.mark.asyncio
async def test_excel_vazifa_holati():
    tasks = [FakeTask("Sport", True, datetime(2026, 8, 2, 9, 0)), FakeTask("O'qish", False)]
    out = await generate_excel_report([], tasks, datetime(2026, 8, 1), datetime(2026, 8, 3))

    sheet = load_workbook(out)["Rejalar"]
    assert sheet.cell(row=2, column=2).value == "Bajarildi"
    assert sheet.cell(row=3, column=2).value == "Bajarilmadi / Kutilmoqda"
    assert sheet.cell(row=3, column=3).value == "Muddat qo'yilmagan"


def test_pdf_yaratiladi():
    out = generate_pdf_report(
        [FakeTx(20000), FakeTx(100000, "income")], [FakeTask("Sport", True)],
        datetime(2026, 8, 1), datetime(2026, 8, 3), user_name="Test",
    )
    assert isinstance(out, io.BytesIO)
    assert out.getvalue().startswith(b"%PDF")
