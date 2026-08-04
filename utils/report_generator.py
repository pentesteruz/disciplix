import io
from datetime import datetime, date
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from database.models import Transaction, Task, Dream, Debt


def _write_sheet(workbook, title: str, rows: list, bosh_xabar: str):
    """
    Lug'atlar ro'yxatini varaqqa yozadi.

    Ilgari bu pandas orqali qilinardi — bitta jadval yozish uchun ~100 MB
    bog'liqlik, holbuki openpyxl (u allaqachon kerak edi) buni o'zi bajaradi.
    """
    sheet = workbook.create_sheet(title=title)

    if not rows:
        sheet.append(["Ma'lumot"])
        sheet.append([bosh_xabar])
        sheet.cell(row=1, column=1).font = Font(bold=True)
        sheet.column_dimensions["A"].width = 40
        return

    headers = list(rows[0].keys())
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        sheet.append([row.get(h, "") for h in headers])

    # Ustun kengligini mazmunga moslash
    for idx, header in enumerate(headers, start=1):
        eng_uzun = max([len(str(header))] + [len(str(r.get(header, ""))) for r in rows])
        sheet.column_dimensions[get_column_letter(idx)].width = min(eng_uzun + 2, 50)


async def generate_excel_report(transactions, tasks, start_date, end_date, lang='uz'):
    """
    Generates an Excel file mapping the given transactions and tasks.
    Returns: io.BytesIO
    """
    workbook = Workbook()
    # Yangi Workbook bo'sh standart varaq bilan keladi — uni olib tashlaymiz.
    workbook.remove(workbook.active)

    _write_sheet(workbook, "Tranzaksiyalar", [
        {
            "Sana": t.timestamp.strftime('%Y-%m-%d %H:%M') if t.timestamp else "",
            "Turi": "Kirim" if t.type == "income" else "Chiqim",
            "Kategoriya": t.category,
            "Miqdor": t.amount,
            "Valyuta": t.currency,
            "Izoh": t.description or "",
        }
        for t in transactions
    ], "Ushbu davrda tranzaksiyalar yo'q")

    _write_sheet(workbook, "Rejalar", [
        {
            "Vazifa": t.title,
            "Holati": "Bajarildi" if t.is_completed else "Bajarilmadi / Kutilmoqda",
            "Muddat": t.due_date.strftime('%Y-%m-%d %H:%M') if t.due_date else "Muddat qo'yilmagan",
        }
        for t in tasks
    ], "Ushbu davrda vazifalar yo'q")

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def generate_pdf_report(transactions, tasks, start_date, end_date, user_name="Foydalanuvchi"):
    """
    Generates a PDF summary of the given interactions.
    Returns: io.BytesIO
    """
    pdf = FPDF()
    pdf.add_page()
    
    # We will use Arial or Helvetica since standard fpdf doesn't support unicode like cyrillic without custom fonts.
    # Usually, Latin characters are fine for Uzbek.
    pdf.set_font("Helvetica", size=16, style='B')
    pdf.cell(200, 10, text=f"{user_name} uchun moliya va rejalar hisoboti", new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    
    pdf.set_font("Helvetica", size=12)
    s_date = start_date.strftime('%Y-%m-%d')
    e_date = end_date.strftime('%Y-%m-%d')
    pdf.cell(200, 10, text=f"Davr: {s_date} dan {e_date} gacha", new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    
    pdf.ln(10)
    
    # Summary of Finance
    total_income = sum(t.amount for t in transactions if t.type == "income")
    total_expense = sum(t.amount for t in transactions if t.type == "expense")
    
    pdf.set_font("Helvetica", size=14, style='B')
    pdf.cell(200, 10, text="Moliyaviy Hisobot (Qisqacha):", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", size=12)
    pdf.cell(200, 10, text=f"Jami kirim (daromad): {total_income:,.0f}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.cell(200, 10, text=f"Jami chiqim (xarajat): {total_expense:,.0f}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    
    pdf.ln(5)
    pdf.set_font("Helvetica", size=14, style='B')
    pdf.cell(200, 10, text="Oxirgi 10 tranzaksiya:", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", size=10)
    
    for t in sorted(transactions, key=lambda x: x.timestamp or datetime.min, reverse=True)[:10]:
        t_type = "Kirim" if t.type == "income" else "Xarajat"
        dt_str = t.timestamp.strftime('%Y-%m-%d') if t.timestamp else ""
        text = f"{dt_str}: {t_type} - {t.amount} {t.currency} ({t.category})"
        # FPDF string handling (ascii/latin-1 limits)
        text = text.encode("latin-1", "replace").decode("latin-1")
        pdf.cell(200, 6, text=text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.ln(10)
    
    # Tasks Summary
    completed = sum(1 for t in tasks if t.is_completed)
    total = len(tasks)
    
    pdf.set_font("Helvetica", size=14, style='B')
    pdf.cell(200, 10, text="Rejalar Hisoboti:", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", size=12)
    pdf.cell(200, 10, text=f"Jami vazifalar: {total}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.cell(200, 10, text=f"Bajarilganlari: {completed}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    
    if total > 0:
        rate = (completed/total)*100
        pdf.cell(200, 10, text=f"Samaradorlik: {rate:.1f}%", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        
    pdf_out = pdf.output()
    if isinstance(pdf_out, str):
        output = io.BytesIO(pdf_out.encode('latin-1'))
    else:
        output = io.BytesIO(bytes(pdf_out))
    output.seek(0)
    
    return output
